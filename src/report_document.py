"""Word typography for the shared, lossless research-report blocks."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def configure_report(document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    section.header_distance = section.footer_distance = Inches(.35)

    styles = document.styles
    for name, size, bold, before, after in (
        ("Normal", 11, False, 0, 8),
        ("Title", 28, True, 8, 20),
        ("Heading 1", 19, True, 24, 12),
        ("Heading 2", 14, True, 18, 8),
        ("Heading 3", 11, True, 12, 6),
        ("Quote", 11, False, 8, 10),
        ("Caption", 9, False, 0, 6),
        ("List Bullet", 11, False, 0, 6),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string("000000" if name != "Caption" else "616B7A")
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        fonts.set(qn("w:eastAsia"), "Heiti SC" if bold else "Songti SC")
        # Some Word templates attach a blue rule to Title by default.
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
        fmt = style.paragraph_format
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.line_spacing = Pt(size * 1.65)
        fmt.widow_control = True
        if name.startswith("Heading") or name == "Title":
            fmt.keep_with_next = True
    styles["Quote"].paragraph_format.left_indent = Inches(.25)
    styles["Quote"].paragraph_format.right_indent = Inches(.25)
    header = section.header.paragraphs[0]
    header.text = "THREADLINE  /  访谈质性分析"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = styles["Caption"]
    footer.add_run("Threadline  ·  ")
    number = OxmlElement("w:fldSimple")
    number.set(qn("w:instr"), "PAGE")
    footer._p.append(number)


def add_report_block(document, kind: str, text: str) -> None:
    levels = {"title": 0, "h1": 1, "h2": 2, "h3": 3}
    if kind in levels:
        document.add_heading(text, levels[kind])
    elif kind == "bullet":
        document.add_paragraph(text, style="List Bullet")
    elif text.startswith("「"):
        document.add_paragraph(text, style="Quote")
    elif text.startswith("（") and text.endswith("）"):
        document.add_paragraph(text, style="Caption")
    else:
        paragraph = document.add_paragraph()
        if text.startswith(("类型：", "来自类别：")):
            paragraph.paragraph_format.keep_with_next = True
        label, separator, body = text.partition("：")
        if separator and len(label) <= 12:
            paragraph.add_run(label + separator).bold = True
            paragraph.add_run(body)
        else:
            paragraph.add_run(text)
