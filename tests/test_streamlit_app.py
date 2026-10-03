import json
import os
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from threading import Event
from types import SimpleNamespace
from unittest.mock import ANY, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from streamlit.proto.TextInput_pb2 import TextInput
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "streamlit_app.py"
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from evidence_fixtures import matched_code
from streamlit_app import api_model_name, build_result_docx, check_api_connection, read_transcript
from api_settings import read_api_key
from analysis_job import JOB_REGISTRY
from analysis_progress import ProgressEvent


class StreamlitAppTests(unittest.TestCase):
    def setUp(self):
        JOB_REGISTRY.clear_completed()
        env_dir = tempfile.TemporaryDirectory()
        self.addCleanup(env_dir.cleanup)
        self.env_path = Path(env_dir.name) / ".env"
        env_patch = patch("api_settings.ENV_PATH", self.env_path)
        env_patch.start()
        self.addCleanup(env_patch.stop)

    def test_uploaded_utf8_text_is_decoded(self):
        uploaded_file = BytesIO("\ufeff 访谈内容 \n".encode("utf-8"))
        self.assertEqual(read_transcript("上传文件", "", uploaded_file), "访谈内容")

    def test_uploaded_docx_reads_paragraphs_and_tables_in_order(self):
        document = Document()
        document.add_paragraph("访谈者：最近有什么变化？")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "受访者"
        table.cell(0, 1).text = "开始远程办公"
        document.add_paragraph("访谈者：谢谢。")

        uploaded_file = BytesIO()
        document.save(uploaded_file)
        uploaded_file.name = "访谈.DOCX"

        self.assertEqual(
            read_transcript("上传文件", "", uploaded_file),
            "访谈者：最近有什么变化？\n受访者\t开始远程办公\n访谈者：谢谢。",
        )

    def test_docx_without_body_text_reports_ocr_requirement(self):
        uploaded_file = BytesIO()
        Document().save(uploaded_file)
        uploaded_file.name = "扫描访谈.docx"

        with self.assertRaisesRegex(ValueError, "需要先进行文字识别"):
            read_transcript("上传文件", "", uploaded_file)

    def test_invalid_docx_reports_read_error(self):
        uploaded_file = BytesIO(b"not a DOCX file")
        uploaded_file.name = "访谈.docx"

        with self.assertRaisesRegex(ValueError, "DOCX 文件无法读取"):
            read_transcript("上传文件", "", uploaded_file)

    def test_docx_with_invalid_body_xml_reports_read_error(self):
        original = BytesIO()
        Document().save(original)
        damaged = BytesIO()
        with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(damaged, "w") as target:
            for item in source.infolist():
                data = b"<broken" if item.filename == "word/document.xml" else source.read(item.filename)
                target.writestr(item, data)
        damaged.name = "访谈.docx"
        with self.assertRaisesRegex(ValueError, "DOCX 文件无法读取"):
            read_transcript("上传文件", "", damaged)

    def test_docx_with_missing_auxiliary_part_recovers_body_text(self):
        document = Document()
        document.add_paragraph("访谈者：最近有什么变化？")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "受访者"
        table.cell(0, 1).text = "开始远程办公"
        original = BytesIO()
        document.save(original)

        damaged = BytesIO()
        with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(damaged, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "_rels/.rels":
                    relationships = ElementTree.fromstring(data)
                    auxiliary = next(
                        rel for rel in relationships
                        if rel.attrib.get("Target", "").endswith("core.xml")
                    )
                    auxiliary.set("Target", "NULL")
                    data = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)
                target.writestr(item, data)

        with self.assertRaises(KeyError):
            Document(BytesIO(damaged.getvalue()))
        damaged.name = "访谈.docx"
        self.assertEqual(
            read_transcript("上传文件", "", damaged),
            "访谈者：最近有什么变化？\n受访者\t开始远程办公",
        )

    def test_docx_without_body_part_reports_clear_error(self):
        original = BytesIO()
        Document().save(original)
        damaged = BytesIO()
        with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(damaged, "w") as target:
            for item in source.infolist():
                if item.filename != "word/document.xml":
                    target.writestr(item, source.read(item.filename))
        damaged.name = "访谈.docx"

        with self.assertRaisesRegex(ValueError, "DOCX 文件结构异常，无法读取正文"):
            read_transcript("上传文件", "", damaged)

    def test_saved_word_report_contains_themes_and_codes(self):
        result = {
            "session_name": "report-test", "accepted": True,
            "metadata": {"final_average_score": 4.25},
            "refinement_iterations": 1,
            "final_evaluation": {
                "global_feedback": "主题整体达标，但有一项置信度偏低。",
                "flagged_themes": ["工作方式变化"],
            },
            "final_themes": [{
                "name": "工作方式变化", "description": "受访者描述了远程办公。",
                "codes": ["通勤时间减少", "线上沟通增加"],
            }],
        }

        document = Document(BytesIO(build_result_docx(result)))
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("访谈质性分析报告", "工作方式变化", "受访者描述了远程办公。",
                         "通勤时间减少", "线上沟通增加", "建议人工复核"):
            self.assertIn(expected, paragraphs)

    def test_saved_word_report_includes_quote_and_evidence_gap(self):
        result = {
            "session_name": "case-01", "accepted": True,
            "metadata": {"final_average_score": 4.0}, "refinement_iterations": 1,
            "codes": [{
                "code_id": 2, "description": "合同由法务掌握",
                "excerpt": "合同在法务部门", "statement_type": "participant_report",
                "verification_status": "reported_only", "open_question": "需向法务核查",
            }],
            "final_themes": [{
                "name": "证据掌握关系", "description": "合同由法务保管。",
                "kind": "evidence_gap", "code_ids": [2],
                "counterexample_code_ids": [], "open_questions": ["公开状态待查"],
                "codes": ["合同由法务掌握"],
            }],
        }

        result["codes"] = [matched_code(**code).model_dump() for code in result["codes"]]
        document = Document(BytesIO(build_result_docx(result)))
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for expected in ("证据缺口", "合同在法务部门", "仅有受访者陈述", "需向法务核查", "公开状态待查"):
            self.assertIn(expected, paragraphs)

    def test_result_page_offers_word_and_json_saves(self):
        result = {
            "session_name": "report-test", "accepted": True,
            "metadata": {"final_average_score": 4.25},
            "refinement_iterations": 1, "final_themes": [],
            "final_evaluation": {
                "global_feedback": "主题已达标，但建议复核。",
                "flagged_themes": ["主题甲"],
            },
        }
        with patch.dict(os.environ, {}, clear=True), patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(APP_PATH))
            app.session_state["analysis_result"] = result
            app.run()

        self.assertFalse(app.exception)
        self.assertEqual([call.args[0] for call in download.call_args_list], [
            "保存 Word 报告（DOCX）", "保存完整数据（JSON）",
        ])
        self.assertEqual(download.call_args_list[0].kwargs["file_name"], "report-test.docx")
        self.assertEqual(json.loads(download.call_args_list[1].kwargs["data"]), result)
        self.assertIn("主题甲", app.warning[0].value)
        app.button(key="clear_analysis_result").click().run()
        self.assertNotIn("analysis_result", app.session_state)

    def test_empty_transcript_shows_validation_error(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.button(key="run_analysis").click().run()

        self.assertFalse(app.exception)
        self.assertIn("请先粘贴访谈文本", app.error[0].value)

    def test_page_shows_methodology_references_at_the_bottom(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()

        self.assertFalse(app.exception)
        headings = [item.value for item in app.header]
        self.assertLess(headings.index("1. 输入访谈材料"), headings.index("2. 运行分析"))
        self.assertLess(headings.index("2. 运行分析"), headings.index("3. 分析结果"))
        self.assertLess(headings.index("3. 分析结果"), headings.index("方法依据与参考文献"))
        references = next(
            item for item in app.expander if item.label == "查看切块策略说明和文献"
        ).markdown[0].value
        self.assertIn("什么让我惊喜？", [item.label for item in app.text_area])
        self.assertIn("TAMA 原始方法", references)
        self.assertIn("Lost in the Middle", references)
        self.assertIn("Token 用量计算", references)

        app.text_area(key="transcript_text").set_value("访谈者：工作有什么变化？\n受访者：开始远程办公。 ").run()
        self.assertIn("切块预估", [item.value for item in app.subheader])
        self.assertTrue(any("预计" in item.value and "片段" in item.value for item in app.markdown))

    def test_reflexivity_notes_save_locally_and_are_not_sent_to_analysis(self):
        with tempfile.TemporaryDirectory() as temp, patch.dict(
            os.environ, {"TAMA_REFLEXIVITY_PATH": str(Path(temp) / "reflexivity.md")}, clear=True,
        ), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.return_value = {
                "session_name": "local-notes", "accepted": True,
                "metadata": {"final_average_score": 4.0},
                "refinement_iterations": 1, "final_themes": [],
            }
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_area(key="reflexivity_surprise").set_value("表格交上去之后没人看").run()
            app.button(key="save_reflexivity").click().run()
            saved = (Path(temp) / "reflexivity.md").read_text(encoding="utf-8")
            self.assertIn("表格交上去之后没人看", saved)
            self.assertIn("不会发送给模型", saved)
            app.text_area(key="transcript_text").set_value("一段测试访谈").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            self.assertNotIn("reflexivity", framework.call_args.kwargs)
            self.assertNotIn("表格交上去之后没人看", framework.return_value.run_analysis.call_args.kwargs["transcript"])

    def test_analysis_displays_theme_without_network_request(self):
        result = {
            "session_name": "test-session",
            "accepted": True,
            "metadata": {"final_average_score": 4.25},
            "refinement_iterations": 1,
            "final_themes": [
                {"name": "治疗焦虑", "description": "受访者描述了治疗前的担忧。", "codes": ["术前担忧"]}
            ],
        }
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.return_value = result
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("MiMo").run()
            app.text_area(key="transcript_text").set_value("一段测试访谈文本").run()
            app.text_input(key="case_id").set_value("Case-01").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            self.assertEqual(app.selectbox(key="chunk_strategy").value, "自动 · 均衡")
            app.selectbox(key="chunk_strategy").set_value("手动设置").run()
            self.assertEqual(app.number_input(key="chunk_size").value, 4000)
            app.number_input(key="chunk_size").set_value(600).run()
            self.assertEqual(app.number_input(key="max_workers").value, 8)
            app.number_input(key="max_workers").set_value(2).run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            app.run()

        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["analysis_result"], result)
        framework.assert_called_once()
        expected = {"api_key": "test-key", "model": "mimo-v2.6-pro",
                    "base_url": "https://api.xiaomimimo.com/v1", "chunk_size": 600,
                    "chunk_strategy": "manual", "max_workers": 2}
        for key, value in expected.items():
            with self.subTest(setting=key):
                self.assertEqual(framework.call_args.kwargs[key], value)
        framework.return_value.run_analysis.assert_called_once_with(
            transcript="一段测试访谈文本",
            save_intermediate=False,
            save_final=False,
            before_model_call=ANY,
            on_progress=ANY,
            concurrency_limit=ANY,
            case_id="Case-01",
        )

    def test_focus_and_local_save_reach_framework(self):
        result = {
            "session_name": "privacy", "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1, "final_themes": [],
        }
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.return_value = result
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="research_profile").set_value("企业能力与证据链").run()
            app.text_input(key="research_question").set_value("信息如何公开？").run()
            app.text_input(key="focus_areas").set_value("证据, 断点").run()
            app.checkbox(key="save_final").set_value(True).run()
            app.text_area(key="transcript_text").set_value("测试访谈").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))

        self.assertEqual(framework.call_args.kwargs["profile"], "enterprise_evidence")
        self.assertEqual(framework.call_args.kwargs["research_question"], "信息如何公开？")
        self.assertEqual(framework.call_args.kwargs["focus_areas"], ["证据", "断点"])
        self.assertTrue(framework.return_value.run_analysis.call_args.kwargs["save_final"])

    def test_default_analysis_settings_reach_framework(self):
        result = {
            "session_name": "mimo-test-session",
            "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1,
            "final_themes": [],
        }
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.return_value = result
            app = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(app.selectbox(key="provider").value, "MiMo")
            self.assertEqual(app.text_input(key="model_MiMo").value, "mimo-v2.6-pro")
            app.text_area(key="transcript_text").set_value("一段测试访谈文本").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            app.run()

        self.assertFalse(app.exception)
        framework.assert_called_once_with(
            api_key="test-key",
            model="mimo-v2.6-pro",
            base_url="https://api.xiaomimimo.com/v1",
            chunk_size=None,
            chunk_strategy="balanced",
            max_workers=8,
            max_iterations=5, corpus_review=True,
            decision_provider=None,
            confidence_threshold=0.7,
            output_dir=str(ROOT / "outputs"),
            profile="generic", research_question="", focus_areas=None,
        )
        framework.return_value.run_analysis.assert_called_once_with(
            transcript="一段测试访谈文本",
            save_intermediate=False,
            save_final=False,
            before_model_call=ANY,
            on_progress=ANY,
            concurrency_limit=ANY,
        )

    def test_api_connection_uses_configured_endpoint_with_one_short_request(self):
        with patch("streamlit_app.OpenAI") as client:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            check_api_connection("test-key", "DeepSeek-V4.1-Flash", "https://api.deepseek.com", "DeepSeek")

        client.assert_called_once_with(
            api_key="test-key", base_url="https://api.deepseek.com",
            timeout=20.0, max_retries=0,
        )
        client.return_value.chat.completions.create.assert_called_once_with(
            model="deepseek-flash",
            messages=[{"role": "user", "content": "请回复 OK。"}],
            max_tokens=128,
        )

    def test_custom_deepseek_model_id_is_preserved(self):
        self.assertEqual(api_model_name("DeepSeek", "deepseek-v4-pro"), "deepseek-v4-pro")

    def test_custom_provider_replaces_openai_with_empty_configuration(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(app.selectbox(key="provider").options, ["MiMo", "DeepSeek", "自定义"])
            app.selectbox(key="provider").set_value("自定义").run()

        self.assertFalse(app.exception)
        self.assertEqual(app.text_input(key="model_自定义").value, "")
        self.assertEqual(app.text_input(key="custom_base_url").value, "")
        self.assertEqual(app.text_input(key="api_key_自定义").proto.type, TextInput.DEFAULT)

    def test_existing_openai_selection_moves_to_custom_provider(self):
        with patch.dict(os.environ, {"OPENAI_API_KEY": "old-test-key"}, clear=True):
            app = AppTest.from_file(str(APP_PATH))
            app.session_state["provider"] = "OpenAI"
            app.session_state["model_OpenAI"] = "gpt-4o"
            app.run()

        self.assertFalse(app.exception)
        self.assertEqual(app.selectbox(key="provider").value, "自定义")
        self.assertEqual(app.text_input(key="model_自定义").value, "")
        self.assertEqual(app.text_input(key="custom_base_url").value, "")

    def test_custom_configuration_reaches_connection_test_and_analysis(self):
        result = {
            "session_name": "custom-test", "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1, "final_themes": [],
        }
        with patch.dict(os.environ, {}, clear=True), patch("openai.OpenAI") as client, \
                patch("tama.TAMAFramework") as framework:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            framework.return_value.run_analysis.return_value = result
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("自定义").run()
            app.text_input(key="api_key_自定义").set_value(" custom-test-key ").run()
            app.text_input(key="model_自定义").set_value(" DeepSeek-V4.1-Flash ").run()
            app.text_input(key="custom_base_url").set_value(" https://custom.example/v1 ").run()
            app.button(key="test_api_connection").click().run()
            self.assertIn("连接成功", app.success[0].value)
            app.text_area(key="transcript_text").set_value("测试访谈").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            app.run()

        self.assertFalse(app.exception)
        client.assert_called_once_with(
            api_key="custom-test-key", base_url="https://custom.example/v1",
            timeout=20.0, max_retries=0,
        )
        self.assertEqual(
            client.return_value.chat.completions.create.call_args.kwargs["model"],
            "DeepSeek-V4.1-Flash",
        )
        framework.assert_called_once()
        for field, value in {
            "api_key": "custom-test-key", "model": "DeepSeek-V4.1-Flash",
            "base_url": "https://custom.example/v1",
        }.items():
            self.assertEqual(framework.call_args.kwargs[field], value)
        self.assertEqual(app.session_state["analysis_result"], result)

    def test_custom_provider_reads_its_own_environment_configuration(self):
        with patch.dict(os.environ, {
            "CUSTOM_API_KEY": "custom-env-test", "CUSTOM_MODEL": "vendor-model",
            "CUSTOM_BASE_URL": "https://custom.example/v1", "OPENAI_API_KEY": "old-test-key",
        }, clear=True), patch("openai.OpenAI") as client:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("自定义").run()
            self.assertEqual(app.text_input(key="model_自定义").value, "vendor-model")
            self.assertEqual(app.text_input(key="custom_base_url").value, "https://custom.example/v1")
            app.button(key="test_api_connection").click().run()

        self.assertFalse(app.exception)
        client.assert_called_once_with(
            api_key="custom-env-test", base_url="https://custom.example/v1",
            timeout=20.0, max_retries=0,
        )

    def test_custom_provider_missing_fields_block_requests(self):
        fields = {
            "api_key_自定义": ("custom-test", "请填写 API Key"),
            "model_自定义": ("vendor-model", "请填写模型名称"),
            "custom_base_url": ("https://custom.example/v1", "请填写接口地址"),
        }
        for missing, (_, message) in fields.items():
            with self.subTest(field=missing), patch.dict(os.environ, {}, clear=True), \
                    patch("openai.OpenAI") as client, patch("tama.TAMAFramework") as framework:
                app = AppTest.from_file(str(APP_PATH)).run()
                app.selectbox(key="provider").set_value("自定义").run()
                for field, (value, _) in fields.items():
                    app.text_input(key=field).set_value(" " if field == missing else value)
                app.run()
                app.button(key="test_api_connection").click().run()
                self.assertIn(message, app.error[0].value)
                app.text_area(key="transcript_text").set_value("测试访谈").run()
                app.button(key="run_analysis").click().run()
                self.assertFalse(app.exception)
                self.assertIn(message, app.error[0].value)
                client.assert_not_called()
                framework.assert_not_called()

    def test_custom_provider_key_is_saved_separately_and_reused(self):
        self.env_path.write_text('DEEPSEEK_API_KEY="deepseek-test"\n', encoding="utf-8")
        with patch.dict(os.environ, {
                    "CUSTOM_MODEL": "vendor-model", "CUSTOM_BASE_URL": "https://custom.example/v1",
                }, clear=True), patch("openai.OpenAI") as client:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("自定义").run()
            app.text_input(key="api_key_自定义").set_value(" saved-custom-test ").run()
            app.button(key="save_api_key").click().run()
            self.assertEqual(read_api_key("自定义"), "saved-custom-test")
            self.assertEqual(read_api_key("DeepSeek"), "deepseek-test")
            self.assertEqual(app.text_input(key="api_key_自定义").value, "saved-custom-test")
            new_session = AppTest.from_file(str(APP_PATH)).run()
            new_session.selectbox(key="provider").set_value("自定义").run()
            new_session.button(key="test_api_connection").click().run()

        self.assertFalse(new_session.exception)
        client.assert_called_once_with(
            api_key="saved-custom-test", base_url="https://custom.example/v1",
            timeout=20.0, max_retries=0,
        )

    def test_existing_deepseek_default_is_relabelled_in_open_session(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH))
            app.session_state["provider"] = "DeepSeek"
            app.session_state["model_DeepSeek"] = "deepseek-flash"
            app.run()

        self.assertFalse(app.exception)
        self.assertEqual(app.text_input(key="model_DeepSeek").value, "DeepSeek-V4.1-Flash")

    def test_mimo_model_migration_preserves_other_model_choices(self):
        for previous, expected in (("mimo-v2.5-pro", "mimo-v2.6-pro"),
                                   ("mimo-v2.5", "mimo-v2.6-pro"),
                                   ("mimo-v2.6-flash", "mimo-v2.6-flash")):
            with self.subTest(previous=previous), patch.dict(os.environ, {}, clear=True):
                app = AppTest.from_file(str(APP_PATH))
                app.session_state["provider"] = "MiMo"
                app.session_state["model_MiMo"] = previous
                app.run()
                self.assertFalse(app.exception)
                self.assertEqual(app.text_input(key="model_MiMo").value, expected)

    def test_mimo_connection_uses_documented_completion_limit(self):
        with patch("streamlit_app.OpenAI") as client:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            check_api_connection("test-key", "mimo-v2.6-pro", "https://api.xiaomimimo.com/v1", "MiMo")

        self.assertEqual(
            client.return_value.chat.completions.create.call_args.kwargs["max_completion_tokens"],
            256,
        )

    def test_pause_and_resume_buttons_control_running_analysis(self):
        first_started = Event()
        release_first = Event()
        second_started = Event()
        result = {
            "session_name": "paused-test", "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1, "final_themes": [],
        }

        def run_analysis(*, before_model_call, **kwargs):
            before_model_call("提取编码")
            first_started.set()
            if not release_first.wait(10):
                raise TimeoutError("first model call timed out")
            before_model_call("归纳主题")
            second_started.set()
            return result

        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.side_effect = run_analysis
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("MiMo").run()
            app.text_area(key="transcript_text").set_value("测试文本").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            try:
                app.button(key="run_analysis").click().run(timeout=6)
                self.assertTrue(first_started.wait(2))
                self.assertTrue(app.selectbox(key="provider").disabled)
                self.assertTrue(app.radio[0].disabled)
                self.assertTrue(app.text_area(key="transcript_text").disabled)
                self.assertTrue(app.text_input(key="model_MiMo").disabled)
                self.assertFalse(app.number_input(key="max_workers").disabled)
                app.button(key="pause_analysis").click().run()
                self.assertTrue(app.session_state["analysis_job"].snapshot().pause_requested)
                release_first.set()
                self.assertFalse(second_started.wait(0.05))
                app.button(key="resume_analysis").click().run()
                self.assertTrue(app.session_state["analysis_job"].done.wait(2))
                app.run()
            finally:
                release_first.set()
                if first_started.is_set():
                    app.session_state["analysis_job"].resume()

        self.assertFalse(app.exception)
        self.assertTrue(second_started.is_set())
        self.assertEqual(app.session_state["analysis_result"], result)

    def test_running_progress_and_concurrency_controls_use_actual_job(self):
        entered, release = Event(), Event()

        def run_analysis(*, before_model_call, on_progress, concurrency_limit, **kwargs):
            before_model_call("提取编码")
            on_progress(ProgressEvent("phase", "提取编码", "coding", total=12, unit="片段"))
            for index in range(1, 5):
                on_progress(ProgressEvent("started", "提取编码", "coding", index=index))
            on_progress(ProgressEvent("completed", "提取编码", "coding", index=3))
            entered.set()
            release.wait(15)
            return {"session_name": "progress-test", "accepted": True, "metadata": {}, "final_themes": []}

        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.side_effect = run_analysis
            app = AppTest.from_file(str(APP_PATH)).run()
            app.button(key="concurrency_preset_16").click().run()
            self.assertEqual(app.number_input(key="max_workers").value, 16)
            app.text_area(key="transcript_text").set_value("模拟材料").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            try:
                app.button(key="run_analysis").click().run()
                self.assertTrue(entered.wait(2))
                app.run()
                self.assertFalse(app.exception)
                self.assertEqual(app.get("progress")[0].proto.value, 8)
                self.assertTrue(any(item.value == "3 / 16" for item in app.metric))
                self.assertIn("已处理 1 / 12", app.get("progress")[0].proto.text)
                app.slider(key="live_max_workers").set_value(10).run()
                job = app.session_state["analysis_job"]
                self.assertEqual(job.concurrency(), 10)
                self.assertEqual(app.number_input(key="max_workers").value, 10)
                app.button(key="concurrency_preset_8").click().run()
                self.assertEqual(job.concurrency(), 8)
                refreshed = AppTest.from_file(str(APP_PATH)).run()
                self.assertEqual(refreshed.slider(key="live_max_workers").value, 8)
                self.assertTrue(any(item.value == "3 / 8" for item in refreshed.metric))
            finally:
                release.set()
                if entered.is_set():
                    self.assertTrue(app.session_state["analysis_job"].done.wait(2))

    def test_legacy_runner_does_not_offer_live_concurrency(self):
        entered, release = Event(), Event()

        def run(checkpoint):
            checkpoint("归并编码")
            entered.set()
            release.wait(10)
            return {}

        with patch.dict(os.environ, {}, clear=True):
            job = JOB_REGISTRY.start(run, max_workers=4)
            try:
                self.assertTrue(entered.wait(2))
                app = AppTest.from_file(str(APP_PATH)).run()
                self.assertFalse(app.exception)
                self.assertTrue(app.number_input(key="max_workers").disabled)
                self.assertNotIn("live_max_workers", [slider.key for slider in app.slider])
                self.assertTrue(any("沿用启动时" in item.value for item in app.caption))
            finally:
                release.set()
                self.assertTrue(job.done.wait(2))

    def test_cleared_concurrency_restores_last_setting(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.number_input(key="max_workers").set_value(12).run()
            app.number_input(key="max_workers").set_value(None).run()
        self.assertFalse(app.exception)
        self.assertEqual(app.number_input(key="max_workers").value, 12)

    def test_new_browser_session_reconnects_to_running_job_and_result(self):
        entered = Event()
        release = Event()
        result = {
            "session_name": "reconnected", "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1, "final_themes": [],
        }

        def run_analysis(*, before_model_call, **kwargs):
            before_model_call("提取编码")
            entered.set()
            if not release.wait(5):
                raise TimeoutError("model call timed out")
            return result

        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.side_effect = run_analysis
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_area(key="transcript_text").set_value("测试文本").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            try:
                app.button(key="run_analysis").click().run()
                self.assertTrue(entered.wait(2))
                refreshed = AppTest.from_file(str(APP_PATH)).run()
                self.assertIs(refreshed.session_state["analysis_job"], app.session_state["analysis_job"])
                self.assertTrue(refreshed.button(key="run_analysis").disabled)
                release.set()
                self.assertTrue(app.session_state["analysis_job"].done.wait(2))
                refreshed.run()
                self.assertEqual(refreshed.session_state["analysis_result"], result)
            finally:
                release.set()

    def test_cancel_shows_pending_state_then_confirmed_cancellation(self):
        entered, release = Event(), Event()

        def run(checkpoint):
            checkpoint("提取编码")
            entered.set()
            release.wait(10)
            return {"session_name": "late-result", "accepted": True,
                    "metadata": {}, "final_themes": []}

        with patch.dict(os.environ, {}, clear=True):
            job = JOB_REGISTRY.start(run)
            try:
                self.assertTrue(entered.wait(2))
                app = AppTest.from_file(str(APP_PATH)).run()
                app.button(key="cancel_analysis").click().run()
                self.assertFalse(app.exception)
                status = app.get("status")[0]
                self.assertEqual(status.proto.label, "正在取消分析")
                self.assertEqual(status.proto.icon, "spinner")
                self.assertTrue(app.button(key="cancel_analysis").disabled)
                self.assertEqual(app.button(key="cancel_analysis").label, "正在取消…")
                self.assertTrue(app.button(key="run_analysis").disabled)
                self.assertNotIn("pause_analysis", [button.key for button in app.button])
                self.assertNotIn("resume_analysis", [button.key for button in app.button])
                self.assertFalse(any("正在分析" in item.value for item in app.info))
                release.set()
                self.assertTrue(job.done.wait(2))
                app.run()
                self.assertFalse(app.exception)
                self.assertEqual(app.get("status")[0].proto.label, "分析已取消")
                self.assertNotEqual(app.get("status")[0].proto.icon, "spinner")
                self.assertFalse(app.button(key="run_analysis").disabled)
                self.assertNotIn("analysis_result", app.session_state)
            finally:
                release.set()
                self.assertTrue(job.done.wait(2))

    def test_new_browser_session_shows_pending_cancellation(self):
        entered, release = Event(), Event()

        def run(checkpoint):
            checkpoint("提取编码")
            entered.set()
            release.wait(10)
            checkpoint("归纳主题")
            return {}

        with patch.dict(os.environ, {}, clear=True):
            job = JOB_REGISTRY.start(run)
            try:
                self.assertTrue(entered.wait(2))
                job.pause()
                job.cancel()
                app = AppTest.from_file(str(APP_PATH)).run()
                self.assertFalse(app.exception)
                self.assertEqual(app.get("status")[0].proto.label, "正在取消分析")
                self.assertTrue(app.button(key="cancel_analysis").disabled)
                self.assertNotIn("resume_analysis", [button.key for button in app.button])
            finally:
                release.set()
                self.assertTrue(job.done.wait(2))


    def test_connection_button_uses_selected_provider_without_network(self):
        with patch.dict(os.environ, {}, clear=True), patch("openai.OpenAI") as client:
            client.return_value.chat.completions.create.return_value.choices = [object()]
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("DeepSeek").run()
            app.text_input(key="api_key_DeepSeek").set_value("test-key").run()
            app.button(key="test_api_connection").click().run()

        self.assertFalse(app.exception)
        self.assertIn("连接成功", app.success[0].value)
        client.assert_called_once_with(
            api_key="test-key", base_url="https://api.deepseek.com",
            timeout=20.0, max_retries=0,
        )
        self.assertEqual(
            client.return_value.chat.completions.create.call_args.kwargs["model"],
            "deepseek-flash",
        )

    def test_loaded_key_is_visible_and_visibility_preserves_edits(self):
        self.env_path.write_text('MIMO_API_KEY="stored-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {"MIMO_API_KEY": "env-test-key"}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "stored-test-key")
            self.assertEqual(app.text_input(key="api_key_MiMo").proto.type, TextInput.DEFAULT)
            self.assertEqual(app.button(key="toggle_api_key_MiMo").label, "隐藏 Key")
            app.button(key="toggle_api_key_MiMo").click().run()
            self.assertEqual(app.text_input(key="api_key_MiMo").proto.type, TextInput.PASSWORD)
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "stored-test-key")
            app.text_input(key="api_key_MiMo").set_value("edited-test-key").run()
            app.button(key="toggle_api_key_MiMo").click().run()
            self.assertEqual(app.text_input(key="api_key_MiMo").proto.type, TextInput.DEFAULT)
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "edited-test-key")
            app.selectbox(key="provider").set_value("DeepSeek").run()
            self.assertEqual(app.text_input(key="api_key_DeepSeek").value, "")
            app.selectbox(key="provider").set_value("MiMo").run()
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "edited-test-key")

        self.assertFalse(app.exception)

    def test_manual_env_edits_refresh_input_and_preserve_unsaved_edits(self):
        self.env_path.write_text('MIMO_API_KEY="stored-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            self.env_path.write_text('MIMO_API_KEY="file-edited-test"\n', encoding="utf-8")
            app.run()
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "file-edited-test")
            app.text_input(key="api_key_MiMo").set_value("unsaved-test").run()
            self.env_path.write_text('MIMO_API_KEY="file-updated-again"\n', encoding="utf-8")
            app.run()
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "unsaved-test")

        self.assertFalse(app.exception)

    def test_existing_empty_widget_displays_previously_saved_key(self):
        self.env_path.write_text('MIMO_API_KEY="stored-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH))
            app.session_state["provider"] = "MiMo"
            app.session_state["api_key_MiMo"] = ""
            app.run()

        self.assertFalse(app.exception)
        self.assertEqual(app.text_input(key="api_key_MiMo").value, "stored-test-key")

    def test_save_replaces_old_key_and_keeps_new_key_in_input(self):
        self.env_path.write_text(
            'MIMO_API_KEY="old-test-key"\nDEEPSEEK_API_KEY="other-test-key"\n', encoding="utf-8",
        )
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_input(key="api_key_MiMo").set_value(" new-test-key ").run()
            app.button(key="save_api_key").click().run()
            self.assertEqual(read_api_key("MiMo"), "new-test-key")
            self.assertEqual(read_api_key("DeepSeek"), "other-test-key")
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "new-test-key")
            new_session = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(new_session.text_input(key="api_key_MiMo").value, "new-test-key")

        self.assertFalse(app.exception)
        self.assertFalse(new_session.exception)

    def test_cleared_input_blocks_requests_without_saved_or_environment_fallback(self):
        self.env_path.write_text('MIMO_API_KEY="stored-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {"MIMO_API_KEY": "env-test-key"}, clear=True), \
                patch("openai.OpenAI") as client, patch("tama.TAMAFramework") as framework:
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_input(key="api_key_MiMo").set_value("").run()
            app.button(key="test_api_connection").click().run()
            self.assertIn("请填写 API Key", app.error[0].value)
            app.text_area(key="transcript_text").set_value("测试访谈").run()
            app.button(key="run_analysis").click().run()
            self.assertIn("请填写 API Key", app.error[0].value)
            client.assert_not_called()
            framework.assert_not_called()
            app.selectbox(key="provider").set_value("DeepSeek").run()
            app.selectbox(key="provider").set_value("MiMo").run()
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "")
            app.button(key="save_api_key").click().run()
            self.assertEqual(read_api_key("MiMo"), "")
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "")
            app.button(key="test_api_connection").click().run()
            self.assertIn("请填写 API Key", app.error[0].value)
            client.assert_not_called()
            new_session = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(new_session.text_input(key="api_key_MiMo").value, "")

        self.assertFalse(app.exception)

    def test_saving_empty_key_without_previous_key_succeeds(self):
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.button(key="save_api_key").click().run()

        self.assertFalse(app.exception)
        self.assertIn("已清除", app.success[0].value)
        self.assertTrue(self.env_path.exists())
        self.assertEqual(read_api_key("MiMo"), "")

    def test_clear_failure_is_reported_without_reusing_old_key(self):
        self.env_path.write_text('MIMO_API_KEY="stored-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {}, clear=True), patch(
            "api_settings.set_key", side_effect=OSError("private-error-detail"),
        ), patch("openai.OpenAI") as client:
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_input(key="api_key_MiMo").set_value("").run()
            app.button(key="save_api_key").click().run()
            self.assertIn("保存失败", app.error[0].value)
            self.assertNotIn("private-error-detail", app.error[0].value)
            app.button(key="test_api_connection").click().run()
            self.assertIn("请填写 API Key", app.error[0].value)
            client.assert_not_called()

        self.assertFalse(app.exception)

    def test_jev_key_uses_same_visibility_and_clear_on_save_behavior(self):
        self.env_path.write_text('JEV_API_KEY="jev-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="decision_mode").set_value("Jev（实验性）").run()
            self.assertEqual(app.text_input(key="api_key_Jev").value, "jev-test-key")
            self.assertEqual(app.text_input(key="api_key_Jev").proto.type, TextInput.DEFAULT)
            app.button(key="toggle_api_key_Jev").click().run()
            self.assertEqual(app.text_input(key="api_key_Jev").proto.type, TextInput.PASSWORD)
            app.text_input(key="api_key_Jev").set_value("").run()
            app.button(key="save_jev_api_key").click().run()
            self.assertEqual(read_api_key("Jev"), "")
            self.assertEqual(app.text_input(key="api_key_Jev").value, "")

        self.assertFalse(app.exception)

    def test_saved_key_is_used_after_new_browser_session(self):
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.return_value = {
                "session_name": "saved-key-test", "accepted": True,
                "metadata": {"final_average_score": 4.0},
                "refinement_iterations": 1, "final_themes": [],
            }
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_input(key="api_key_MiMo").set_value(" saved-test-key ").run()
            app.button(key="save_api_key").click().run()
            self.assertEqual(read_api_key("MiMo"), "saved-test-key")
            self.assertEqual(app.text_input(key="api_key_MiMo").value, "saved-test-key")

            new_session = AppTest.from_file(str(APP_PATH)).run()
            self.assertEqual(new_session.text_input(key="api_key_MiMo").value, "saved-test-key")
            new_session.text_area(key="transcript_text").set_value("测试访谈").run()
            new_session.button(key="run_analysis").click().run()
            self.assertTrue(new_session.session_state["analysis_job"].done.wait(2))
            new_session.run()

        self.assertFalse(new_session.exception)
        framework.assert_called_once()
        self.assertEqual(framework.call_args.kwargs["api_key"], "saved-test-key")

    def test_delete_saved_key_removes_it_from_next_session(self):
        self.env_path.write_text('MIMO_API_KEY="saved-test-key"\n', encoding="utf-8")
        with patch.dict(os.environ, {}, clear=True):
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("MiMo").run()
            app.text_input(key="api_key_MiMo").set_value(" ").run()
            app.button(key="save_api_key").click().run()
            self.assertEqual(read_api_key("MiMo"), "")

            new_session = AppTest.from_file(str(APP_PATH)).run()
            new_session.selectbox(key="provider").set_value("MiMo").run()
            new_session.text_area(key="transcript_text").set_value("测试访谈").run()
            new_session.button(key="run_analysis").click().run()

        self.assertFalse(new_session.exception)
        self.assertIn("请填写 API Key", new_session.error[0].value)

    def test_env_failure_allows_manual_key_without_exposing_error(self):
        with patch.dict(os.environ, {}, clear=True), \
                patch("api_settings.dotenv_values", side_effect=OSError("private-error-detail")), \
                patch("api_settings.write_api_key", side_effect=OSError("private-error-detail")):
            app = AppTest.from_file(str(APP_PATH)).run()
            self.assertIn("项目 .env 暂时无法读取", app.warning[0].value)
            self.assertNotIn("private-error-detail", app.warning[0].value)
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="save_api_key").click().run()

        self.assertFalse(app.exception)
        self.assertIn("保存失败", app.error[0].value)
        self.assertNotIn("private-error-detail", app.error[0].value)

    def test_connection_error_does_not_show_raw_exception(self):
        with patch.dict(os.environ, {}, clear=True), patch("openai.OpenAI") as client:
            client.return_value.chat.completions.create.side_effect = RuntimeError("test-key")
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="provider").set_value("MiMo").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="test_api_connection").click().run()

        self.assertFalse(app.exception)
        self.assertIn("连接失败", app.error[0].value)
        self.assertNotIn("test-key", app.error[0].value)

    def test_analysis_type_error_shows_safe_diagnostics(self):
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            framework.return_value.run_analysis.side_effect = TypeError("secret-api-key")
            app = AppTest.from_file(str(APP_PATH)).run()
            app.text_area(key="transcript_text").set_value("测试文本").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            app.run()

        captions = "\n".join(item.value for item in app.caption)
        self.assertIn("错误类型：TypeError", captions)
        self.assertIn("错误位置：", captions)
        self.assertNotIn("secret-api-key", captions)

    def test_jev_mode_without_key_blocks_analysis(self):
        with patch.dict(os.environ, {}, clear=True), patch("tama.TAMAFramework") as framework:
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="decision_mode").set_value("Jev（实验性）").run()
            app.text_area(key="transcript_text").set_value("测试访谈").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()

        self.assertFalse(app.exception)
        self.assertIn("Jev API Key", app.error[0].value)
        framework.assert_not_called()

    def test_jev_mode_builds_jev_decision_client(self):
        result = {
            "session_name": "jev-test", "accepted": True,
            "metadata": {"final_average_score": 4.0},
            "refinement_iterations": 1, "final_themes": [],
        }
        with patch.dict(os.environ, {}, clear=True), \
                patch("tama.TAMAFramework") as framework, \
                patch("decisions.jev_client.JevDecisionClient") as jev_client_cls:
            framework.return_value.run_analysis.return_value = result
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="decision_mode").set_value("Jev（实验性）").run()
            app.text_input(key="api_key_Jev").set_value("jev-test").run()
            app.text_area(key="transcript_text").set_value("测试访谈").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="run_analysis").click().run()
            self.assertTrue(app.session_state["analysis_job"].done.wait(2))
            app.run()

        self.assertFalse(app.exception)
        jev_client_cls.assert_called_once_with(api_key="jev-test", base_url=None, model=None)
        framework.assert_called_once()
        self.assertEqual(framework.call_args.kwargs["api_key"], "test-key")
        self.assertIs(framework.call_args.kwargs["decision_provider"], jev_client_cls.return_value)

    def test_jev_connection_test_sends_noul_request(self):
        with patch.dict(os.environ, {}, clear=True), \
                patch("openai.OpenAI") as main_client, \
                patch("decisions.jev_client.JevDecisionClient") as jev_client_cls:
            main_client.return_value.chat.completions.create.return_value.choices = [object()]
            jev_client_cls.return_value.ask.return_value = {
                "check": SimpleNamespace(probability=0.99),
            }
            app = AppTest.from_file(str(APP_PATH)).run()
            app.selectbox(key="decision_mode").set_value("Jev（实验性）").run()
            app.text_input(key="api_key_Jev").set_value("jev-test").run()
            app.text_input(key="api_key_MiMo").set_value("test-key").run()
            app.button(key="test_api_connection").click().run()

        self.assertFalse(app.exception)
        self.assertIn("连接成功", app.success[0].value)
        jev_client_cls.assert_called_once_with(
            api_key="jev-test", base_url=None, timeout=20.0, max_retries=0,
        )
        state, questions = jev_client_cls.return_value.ask.call_args.args
        self.assertEqual(state, "连接测试")
        self.assertEqual(questions["check"].kind, "noul")


if __name__ == "__main__":
    unittest.main()
