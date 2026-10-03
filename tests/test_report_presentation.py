"""Presentation regressions using synthetic, locally matched evidence only."""

import os
import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from analysis_job import JOB_REGISTRY
from evidence_fixtures import matched_code
from reporting import report_lines
from streamlit_app import build_result_docx
from ui_design import evidence_quote, narrative, theme_outline


def presentation_result():
    return {
        "session_name": "presentation-test", "accepted": False,
        "metadata": {"final_average_score": None}, "refinement_iterations": 1,
        "analytic_storyline": "分享资料的范围随信任关系变化。",
        "configuration": {"save_final": False},
        "codes": [matched_code(code_id=1, name="逐步分享", description="逐步分享",
                               excerpt="熟悉之后，我才会把资料给他。").model_dump(),
                  matched_code(code_id=2, name="拒绝分享", description="拒绝分享",
                               excerpt="认识很久也不一定能给。").model_dump()],
        "final_themes": [{"name": "信任影响资料分享", "kind": "pattern",
                          "description": "资料分享受关系和权限共同影响。",
                          "code_ids": [1, 2], "counterexample_code_ids": [2],
                          "open_questions": ["核对资料权限"], "rationale": "比较两种分享条件。"}],
    }


class ReportPresentationTests(unittest.TestCase):
    def test_report_views_keep_editing_and_export_separate(self):
        JOB_REGISTRY.clear_completed()
        with patch.dict(os.environ, {}, clear=True), patch("keyring.get_password", return_value=None):
            app = AppTest.from_file(str(ROOT / "streamlit_app.py"))
            app.session_state["analysis_result"] = presentation_result()
            app.run()
        self.assertFalse(app.exception)
        self.assertEqual([tab.label for tab in app.tabs], [
            "分析报告", "报告总览", "主题与证据", "研究复核", "导出与记录", "准备材料", "模型连接", "分析参数",
        ])
        findings = next(tab for tab in app.tabs if tab.label == "主题与证据")
        html = "\n".join(item.value for item in findings.markdown)
        self.assertIn('href="#theme-1"', html)
        self.assertIn('id="theme-1"', html)
        self.assertIn("1 条支持证据", html)
        self.assertIn("1 条有效反证", html)
        self.assertIn("熟悉之后，我才会把资料给他。", html)
        self.assertTrue(any(item.label == "添加研究者诠释" for item in findings.expander))
        self.assertTrue(any(item.value == "未评分" for item in app.metric))
        self.assertIsNone(JOB_REGISTRY.current())

    def test_unmatched_code_is_not_promoted_to_featured_quote(self):
        JOB_REGISTRY.clear_completed()
        result = presentation_result()
        result["codes"][0]["evidence_status"] = "unmatched"
        with patch.dict(os.environ, {}, clear=True), patch("keyring.get_password", return_value=None):
            app = AppTest.from_file(str(ROOT / "streamlit_app.py"))
            app.session_state["analysis_result"] = result
            app.run()
        self.assertFalse(app.exception)
        findings = next(tab for tab in app.tabs if tab.label == "主题与证据")
        # The evidence inspection expander may retain the record; the reading
        # surface must never feature it as validated support.
        main_html = "\n".join(item.value for item in findings.markdown if "tl-badges" in item.value)
        self.assertIn("0 条支持证据", main_html)

    def test_custom_html_escapes_all_material_and_model_text(self):
        payload = '<img src=x onerror="alert(1)"> & **text**'
        with patch("ui_design.st.markdown") as render:
            narrative(payload, payload, payload)
            evidence_quote({"code_id": payload, "excerpt": payload, "source": {"source_id": payload}})
            theme_outline([{"name": payload, "description": payload}], navigation=True)
        for call in render.call_args_list:
            html = call.args[0]
            self.assertNotIn("<img", html)
            self.assertIn("&lt;img", html)

    def test_word_typography_preserves_every_report_block(self):
        result = presentation_result()
        document = Document(BytesIO(build_result_docx(result)))
        self.assertEqual([p.text for p in document.paragraphs], [text for _, text in report_lines(result)])
        quotes = [p for p in document.paragraphs if p.style.name == "Quote"]
        self.assertTrue(quotes)
        self.assertIn("熟悉之后", quotes[0].text)
        self.assertGreater(document.styles["Heading 1"].font.size, document.styles["Heading 2"].font.size)
        self.assertGreater(document.styles["Heading 2"].font.size, document.styles["Normal"].font.size)
        self.assertTrue(document.styles["Heading 1"].paragraph_format.keep_with_next)
        self.assertFalse(document.styles["Title"].element.findall('.//' + qn("w:pBdr")))
        self.assertTrue(document.sections[0].footer._element.findall('.//' + qn("w:fldSimple")))


if __name__ == "__main__":
    unittest.main()
