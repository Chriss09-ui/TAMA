"""Local web interface for TAMA thematic analysis."""

import json
import os
import sys
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import keyring
import streamlit as st
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from openai import APIStatusError, OpenAI


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))

from tama import TAMAFramework
from analysis_job import JOB_REGISTRY
from chunking import ChunkStrategy, plan_chunks
from decisions import DecisionQuestion
from decisions.jev_client import JevDecisionClient
from prompts import (
    API_CONNECTION_TEST_PROMPT,
    JEV_CONNECTION_TEST_INSTRUCTIONS,
    JEV_CONNECTION_TEST_STATE,
)
from model_config import api_model_name
from reflexivity import (
    POSITION_FIELD,
    REFLEXIVITY_CHECKS,
    REFLEXIVITY_FIELDS,
    empty_reflexivity,
    load_reflexivity,
    save_reflexivity,
)
from reporting import report_lines

def reflexivity_path() -> Path:
    override = os.environ.get("TAMA_REFLEXIVITY_PATH", "").strip()
    return Path(override) if override else ROOT / "reflexivity.md"

JEV_DECISION_MODE = "Jev（实验性）"

THEME_KIND_LABELS = {
    "pattern": "案例内模式",
    "information_breakpoint": "信息断点",
    "counterexample": "反例",
    "evidence_gap": "证据缺口",
    "candidate_mechanism": "候选机制",
}
STATEMENT_TYPE_LABELS = {
    "material_fact": "材料或事件支持",
    "participant_report": "受访者报告",
    "participant_interpretation": "受访者判断",
    "researcher_interpretation": "研究者解释",
    "undetermined": "待判定",
}
VERIFICATION_LABELS = {
    "supported_by_material": "有材料支持",
    "reported_only": "仅有受访者陈述",
    "conflicting": "存在冲突",
    "unknown": "尚未核实",
}

CHUNK_STRATEGIES: dict[str, ChunkStrategy] = {
    "自动 · 精细": "fine",
    "自动 · 均衡": "balanced",
    "自动 · 少调用": "economy",
    "手动设置": "manual",
}

KEYRING_SERVICE = "TAMA Qualitative Analysis"


def load_saved_api_key(provider: str) -> str:
    """Read a provider key once per browser session."""
    state_key = f"saved_api_key_{provider}"
    if state_key not in st.session_state:
        try:
            st.session_state[state_key] = keyring.get_password(KEYRING_SERVICE, provider) or ""
        except Exception:
            st.session_state[state_key] = ""
            st.session_state[f"api_key_notice_{provider}"] = (
                "warning", "系统凭据库暂不可用；仍可手动填写 API Key。"
            )
    return st.session_state[state_key]


def save_api_key(provider: str) -> None:
    """Persist a key only after the user clicks Save."""
    typed_key = st.session_state.get(f"api_key_{provider}", "").strip()
    notice_key = f"api_key_notice_{provider}"
    if not typed_key:
        st.session_state[notice_key] = ("error", "请先输入要保存的 API Key。")
        return
    try:
        keyring.set_password(KEYRING_SERVICE, provider, typed_key)
    except Exception:
        st.session_state[notice_key] = ("error", "保存失败，请检查系统凭据库是否可用。")
        return
    st.session_state[f"saved_api_key_{provider}"] = typed_key
    st.session_state[f"api_key_{provider}"] = ""
    st.session_state[notice_key] = ("success", "API Key 已保存到系统凭据库。")


def delete_saved_api_key(provider: str) -> None:
    """Remove only this provider's persisted key."""
    notice_key = f"api_key_notice_{provider}"
    try:
        keyring.delete_password(KEYRING_SERVICE, provider)
    except Exception:
        st.session_state[notice_key] = ("error", "删除失败，请检查系统凭据库是否可用。")
        return
    st.session_state[f"saved_api_key_{provider}"] = ""
    st.session_state[notice_key] = ("success", "已删除系统凭据库中保存的 API Key。")


def check_api_connection(api_key: str, model: str, base_url: str | None, provider: str) -> None:
    """Send one short request using the same endpoint and model as analysis."""
    client = OpenAI(api_key=api_key, base_url=base_url, timeout=20.0, max_retries=0)
    token_limit = {"max_completion_tokens": 256} if provider == "MiMo" else {"max_tokens": 128}
    response = client.chat.completions.create(
        model=api_model_name(provider, model),
        messages=[{"role": "user", "content": API_CONNECTION_TEST_PROMPT}],
        **token_limit,
    )
    if not response.choices:
        raise ValueError("模型返回空结果")


def check_jev_connection(api_key: str, base_url: str | None) -> None:
    """Send one minimal yes/no decision to the Jev System One endpoint."""
    client = JevDecisionClient(api_key=api_key, base_url=base_url, timeout=20.0, max_retries=0)
    answers = client.ask(JEV_CONNECTION_TEST_STATE, {
        "check": DecisionQuestion(
            key="check", kind="noul",
            instructions=JEV_CONNECTION_TEST_INSTRUCTIONS,
        ),
    })
    answer = answers["check"]
    if answer.probability is None or answer.probability < 0.5:
        raise ValueError("模型返回异常结果")


def connection_input_error(
    api_key: str, model: str, base_url: str | None, endpoint: str | None,
    decision_mode: str, jev_api_key: str, env_name: str,
) -> str | None:
    if not api_key:
        return f"请填写 API Key，或设置 {env_name} 环境变量。"
    if not model.strip():
        return "请填写模型名称。"
    if base_url is not None and not endpoint:
        return "请填写接口地址。"
    if decision_mode == JEV_DECISION_MODE and not jev_api_key:
        return "决策模式为 Jev 时，请填写 Jev API Key，或设置 JEV_API_KEY 环境变量。"
    return None


@st.fragment(run_every="1s")
def render_active_job() -> None:
    job = st.session_state["analysis_job"]
    snapshot = job.snapshot()
    if snapshot.done:
        st.rerun()

    if snapshot.pause_requested:
        if st.button("继续分析", key="resume_analysis", use_container_width=True):
            job.resume()
            st.rerun()
    elif st.button("暂停分析", key="pause_analysis", use_container_width=True):
        job.pause()
        st.rerun()
    if st.button("取消分析", key="cancel_analysis", use_container_width=True):
        job.cancel()
        st.rerun()

    snapshot = job.snapshot()
    if snapshot.paused:
        st.warning(f"已阻止后续请求，先前发出的请求可能仍在完成 · {snapshot.stage}")
    elif snapshot.pause_requested:
        st.info(f"等待当前请求完成后暂停 · {snapshot.stage}")
    else:
        st.info(f"正在分析 · {snapshot.stage}")
    st.caption("暂停或取消会在下一次模型请求前生效；刷新页面可接回任务，请保持本地服务运行。")


def extract_docx_body_xml(contents: bytes) -> str:
    """Recover body text when a broken auxiliary DOCX relationship blocks python-docx."""
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def paragraph_text(paragraph) -> str:
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{word}t":
                parts.append(node.text or "")
            elif node.tag == f"{word}tab":
                parts.append("\t")
            elif node.tag in (f"{word}br", f"{word}cr"):
                parts.append("\n")
        return "".join(parts).strip()

    with ZipFile(BytesIO(contents)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    body = root.find(f"{word}body")
    if body is None:
        raise ValueError("DOCX 正文不存在。")

    lines = []
    for block in body:
        if block.tag == f"{word}p":
            text = paragraph_text(block)
            if text:
                lines.append(text)
        elif block.tag == f"{word}tbl":
            for row in block.findall(f"{word}tr"):
                cells = []
                for cell in row.findall(f"{word}tc"):
                    cell_text = "\n".join(
                        text for paragraph in cell.iter(f"{word}p")
                        if (text := paragraph_text(paragraph))
                    )
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    lines.append("\t".join(cells))

    if not lines:
        raise ValueError("DOCX 文件中没有可读取的文字。图片或扫描件需要先进行文字识别。")
    return "\n".join(lines)


def extract_docx_text(contents: bytes) -> str:
    """Extract body paragraphs and table rows in document order."""
    try:
        document = Document(BytesIO(contents))
    except KeyError:
        try:
            return extract_docx_body_xml(contents)
        except (BadZipFile, KeyError, ElementTree.ParseError, ValueError) as exc:
            raise ValueError("DOCX 文件结构异常，无法读取正文。请另存为新的 DOCX 后重试。") from exc
    except (BadZipFile, PackageNotFoundError) as exc:
        raise ValueError("DOCX 文件无法读取，请确认文件未损坏且格式正确。") from exc

    lines = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    lines.append(row_text)
        elif block.text.strip():
            lines.append(block.text.strip())

    if not lines:
        raise ValueError("DOCX 文件中没有可读取的文字。图片或扫描件需要先进行文字识别。")
    return "\n".join(lines)


def read_transcript(source: str, pasted_text: str, uploaded_file) -> str:
    """Return pasted text or text extracted from a TXT or DOCX upload."""
    if source == "粘贴文本":
        return pasted_text.strip()
    if uploaded_file is None:
        return ""
    if Path(getattr(uploaded_file, "name", "")).suffix.lower() == ".docx":
        return extract_docx_text(uploaded_file.getvalue())
    return uploaded_file.getvalue().decode("utf-8-sig").strip()


def build_result_docx(result: dict, reflexivity: dict | None = None, interpretations: list | None = None) -> bytes:
    """Create a readable Word report with quotes, analysis, and blank interpretation."""
    document = Document()
    heading_levels = {"title": 0, "h1": 1, "h2": 2, "h3": 3}
    for kind, text in report_lines(result, reflexivity, interpretations):
        level = heading_levels.get(kind)
        if level is not None:
            document.add_heading(text, level)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def reflexivity_from_state() -> dict:
    """Read the on-page reflexivity notes. They are not part of the analysis request."""
    data = empty_reflexivity()
    for key in data:
        data[key] = st.session_state.get(f"reflexivity_{key}", "")
    return data


def ensure_reflexivity_state() -> None:
    if st.session_state.get("_reflexivity_loaded"):
        return
    path = reflexivity_path()
    stored = load_reflexivity(path) if path.is_file() else empty_reflexivity()
    for key, value in stored.items():
        st.session_state.setdefault(f"reflexivity_{key}", value)
    st.session_state["_reflexivity_loaded"] = True


def save_reflexivity_notes() -> None:
    save_reflexivity(reflexivity_path(), reflexivity_from_state())
    st.session_state["reflexivity_notice"] = "自反记录已写在本机 reflexivity.md，没有发送给模型。"


def render_result(result: dict) -> None:
    """Display the latest analysis without exposing the API key."""
    st.header("3. 分析结果")
    st.caption("最近一次成功运行")

    status = {
        "accepted": "达到评估标准", "no_improvement": "连续评估未提升，已早停",
        "max_iterations": "已达到最大评估轮次",
    }.get(result.get("stop_reason"), "达到评估标准" if result["accepted"] else "分析已结束")
    score = result["metadata"]["final_average_score"]
    themes = result["final_themes"]
    code_by_id = {code.get("code_id"): code for code in result.get("codes", [])}

    st.success(f"分析完成 · {status}")
    score_col, theme_col = st.columns(2)
    score_col.metric("平均分", f"{score:.2f} / 5")
    theme_col.metric("主题数", len(themes))
    st.caption(f"评估轮次：{result['refinement_iterations']}")
    storyline = result.get("generation", {}).get("analytic_storyline")
    if storyline:
        st.write(f"初始主题故事线：{storyline}")
    scores = result.get("score_history") or []
    if scores:
        st.caption("各轮平均分：" + " → ".join(f"{item['average_score']:.2f}" for item in scores))
    st.caption("分数衡量分析结果的呈现质量；候选机制仍需研究者结合案例证据核查。")

    final_evaluation = result.get("final_evaluation") or {}
    if final_evaluation.get("global_feedback"):
        st.info(final_evaluation["global_feedback"])
    flagged_themes = final_evaluation.get("flagged_themes", [])
    if flagged_themes:
        st.warning(
            f"建议人工复核 {len(flagged_themes)} 个主题：{'、'.join(flagged_themes)}"
        )

    st.subheader("保存结果")
    interpretations = [
        st.session_state.get(f"theme_interpretation_{index}", "")
        for index in range(len(themes))
    ]
    st.download_button(
        "保存 Word 报告（DOCX）",
        data=build_result_docx(result, reflexivity_from_state(), interpretations),
        file_name=f"{result['session_name']}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )
    st.download_button(
        "保存完整数据（JSON）",
        data=json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8"),
        file_name=f"{result['session_name']}.json",
        mime="application/json",
        use_container_width=True,
    )
    if st.button("清除服务内结果", key="clear_analysis_result", use_container_width=True):
        JOB_REGISTRY.clear_completed()
        st.session_state.pop("analysis_result", None)
        st.session_state.pop("analysis_job", None)
        st.rerun()
    if result.get("configuration", {}).get("save_final"):
        st.caption(f"已保存到本机：{ROOT / 'outputs' / result['session_name']}")
    elif result.get("configuration", {}).get("save_intermediate"):
        st.caption(f"阶段文件已保存到本机：{ROOT / 'outputs' / result['session_name']}；最终结果仅在服务进程内。")
    else:
        st.caption("结果仅保存在当前服务进程内；服务重启后无法恢复，可按需下载。")

    plan_groups = (result.get("next_data_plan") or {}).get("groups") or {}
    if any(plan_groups.values()):
        st.subheader("下一轮可以收集什么")
        for group_name in ("需补访角色", "需查材料", "需核对信息"):
            entries = plan_groups.get(group_name) or []
            if not entries:
                continue
            st.markdown(f"**{group_name}**")
            for entry in entries:
                st.write(f"• {entry}")

    if result.get("memos"):
        with st.expander("分析备忘录"):
            st.caption("模型写下的是理由和不确定。你补在 Word 报告「诠释」栏里的话不会进入模型。")
            for memo in result["memos"]:
                st.markdown(f"**{memo.get('theme') or '未命名主题'}**")
                st.write(memo.get("rationale") or "模型没有写下理由。")
                if memo.get("uncertain"):
                    st.write(f"不确定：{memo['uncertain']}")
                if memo.get("contradicting_code_ids"):
                    st.write(f"反证编码：{', '.join(map(str, memo['contradicting_code_ids']))}")

    if not themes:
        st.info("本次分析没有生成主题。请检查输入文本和模型返回内容。")
        return

    for index, theme in enumerate(themes, 1):
        with st.container(border=True):
            st.subheader(f"{index:02d} · {theme['name']}")
            st.caption(THEME_KIND_LABELS.get(theme.get("kind"), "待分类"))
            st.write(theme["description"])
            if theme.get("rationale"):
                st.write(f"备忘录：{theme['rationale']}")
            code_ids = theme.get("code_ids") or []
            counter_ids = set(theme.get("counterexample_code_ids") or [])
            if code_ids and code_by_id:
                with st.expander(f"查看证据与编码 · {len(code_ids)} 条"):
                    for code_id in code_ids:
                        code = code_by_id.get(code_id)
                        if not code:
                            continue
                        label = code.get("name") or code["description"]
                        st.markdown(f"**[{code_id}] {label}**")
                        if code.get("definition"):
                            st.caption(f"定义：{code['definition']}")
                        if code.get("excerpt"):
                            st.write(f"原话：{code['excerpt']}")
                        st.caption(
                            f"{STATEMENT_TYPE_LABELS.get(code.get('statement_type'), '待判定')} · "
                            f"{VERIFICATION_LABELS.get(code.get('verification_status'), '尚未核实')}"
                        )
                        if code.get("open_question"):
                            st.write(f"待核查：{code['open_question']}")
            elif theme.get("codes"):
                with st.expander(f"查看关联编码 · {len(theme['codes'])} 条"):
                    for code in theme["codes"]:
                        st.write(f"• {code}")
            if counter_ids:
                st.markdown("**反例与边界**")
                for code_id in theme.get("counterexample_code_ids") or []:
                    code = code_by_id.get(code_id)
                    if not code:
                        st.write(f"反例编码：{code_id}")
                        continue
                    st.write(f"[{code_id}] {code.get('name') or code.get('description')}")
                    if code.get("excerpt"):
                        st.write(f"原话：{code['excerpt']}")
            for question in theme.get("open_questions", []):
                st.write(f"待核查：{question}")
            st.text_area(
                "诠释（由你写，不发送给模型）",
                key=f"theme_interpretation_{index - 1}",
                height=80,
                placeholder="所以这一模式意味着什么？",
            )

    canonical_codes = [
        code for code in result.get("codes") or []
        if code.get("merged_into") is None
    ]
    if canonical_codes:
        with st.expander(f"编码簿 · {len(canonical_codes)} 条主编码"):
            for code in canonical_codes:
                label = code.get("name") or code.get("description")
                st.markdown(f"**[{code.get('code_id', '?')}] {label}**")
                if code.get("definition"):
                    st.write(f"定义：{code['definition']}")
                if code.get("include"):
                    st.write(f"包含：{code['include']}")
                if code.get("exclude"):
                    st.write(f"排除：{code['exclude']}")
                if code.get("merged_from"):
                    st.caption(f"并入来源：{', '.join(map(str, code['merged_from']))}")
                if code.get("excerpt"):
                    st.write(f"正例：{code['excerpt']}")
                if code.get("open_question"):
                    st.write(f"待核查：{code['open_question']}")


def render_reflexivity() -> None:
    """Local notes for the researcher. Nothing here is sent with the transcript."""
    ensure_reflexivity_state()
    with st.expander("研究者自反（只保存在本机，不发送给模型）"):
        st.caption("跑完一轮后可以顺手填。这些话进入 Word 报告的立场和自反两节，不进入编码或主题。")
        _key, label, help_text = POSITION_FIELD
        st.text_area(label, key="reflexivity_position", height=80, help=help_text)
        for key, label, help_text in REFLEXIVITY_FIELDS:
            st.text_area(label, key=f"reflexivity_{key}", height=80, help=help_text)
        st.markdown("**定稿前可以自问**")
        for key, label in REFLEXIVITY_CHECKS:
            st.text_area(label, key=f"reflexivity_{key}", height=68)
        st.button("保存自反记录", key="save_reflexivity", on_click=save_reflexivity_notes)
        notice = st.session_state.pop("reflexivity_notice", None)
        if notice:
            st.success(notice)


def render_methodology_references() -> None:
    """Show the methodological sources behind chunking and coding."""
    st.divider()
    st.header("方法依据与参考文献")
    st.caption(
        "文献支持的是“保留语义语境、控制上下文长度、再从片段中识别意义单元”等原则；"
        "1200、1800、2600 等切块数值是本项目的工程参数，不是质性研究的统一标准。"
    )
    with st.expander("查看切块策略说明和文献"):
        st.markdown(
            """
- **TAMA 原始方法**：将过长访谈分成较小片段，并尽量按段落、换行和访谈者标签等自然边界保留对话连贯性。[Xu et al. (2025), *TAMA: A Human-AI Collaborative Thematic Analysis Framework*](https://arxiv.org/abs/2503.20666)

- **意义单元与编码**：文本片段可包含多个意义单元，编码应围绕相对完整的含义，避免单元过宽或过窄。[Graneheim & Lundman (2004), *Qualitative content analysis in nursing research*](https://pubmed.ncbi.nlm.nih.gov/14769454/)

- **编码粒度**：有意义的编码文本通常短于一个段落、长于几个词，编码应表达完整思想。[Saunders et al. (2023), *Practical thematic analysis*](https://www.bmj.com/content/381/bmj-2022-074256)

- **主题分析过程**：编码应覆盖整个数据集，后续主题需同时对照已编码摘录和完整数据集进行审查。[Braun & Clarke, *Doing Reflexive Thematic Analysis*](https://www.thematicanalysis.net/doing-reflexive-ta/)

- **长上下文限制**：模型虽然可以接收很长的上下文，但对中间位置信息的利用可能明显下降。[Liu et al. (2023), *Lost in the Middle*](https://arxiv.org/abs/2307.03172)

- **Token 长度估算**：DeepSeek 文档给出中英文字符到 token 的经验换算；本项目在此基础上增加余量，以兼容不同模型的分词差异。[DeepSeek API 文档：Token 用量计算](https://api-docs.deepseek.com/zh-cn/quick_start/token_usage/)
"""
        )


def main() -> None:
    st.set_page_config(page_title="TAMA · 访谈质性分析", page_icon="📝", layout="wide")
    st.markdown(
        """
        <style>
        .block-container { max-width: 960px; padding-top: 2rem; padding-bottom: 4rem; }
        @media (max-width: 640px) {
            .block-container { padding-top: 1rem; }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    job = st.session_state.get("analysis_job") or JOB_REGISTRY.current()
    if job is not None:
        st.session_state["analysis_job"] = job
    running = job is not None and not job.snapshot().done

    with st.sidebar:
        st.header("分析设置")
        st.subheader("模型连接")
        providers = {
            "DeepSeek": ("DEEPSEEK_API_KEY", "DeepSeek-V4.1-Flash"),
            "MiMo": ("MIMO_API_KEY", "mimo-v2.5-pro"),
            "OpenAI": ("OPENAI_API_KEY", "gpt-4o"),
        }
        provider = st.selectbox("服务商", list(providers), key="provider", disabled=running)
        env_name, default_model = providers[provider]
        saved_api_key = load_saved_api_key(provider)
        api_key_input = st.text_input(
            "API Key", type="password", key=f"api_key_{provider}", disabled=running,
        )
        if api_key_input.strip():
            st.caption("本次使用输入框中的 API Key。点击保存后，下次打开仍可使用。")
        elif saved_api_key:
            st.caption("凭据来源：本机系统凭据库。输入框留空时自动使用。")
        elif os.getenv(env_name, "").strip():
            st.caption(f"凭据来源：环境变量 {env_name}。")
        else:
            st.caption(f"尚未配置。可输入密钥，或设置 {env_name} 环境变量。")
        save_col, delete_col = st.columns(2)
        with save_col:
            st.button(
                "保存 API Key", key="save_api_key", on_click=save_api_key,
                args=(provider,), disabled=running, use_container_width=True,
            )
        with delete_col:
            st.button(
                "删除已存 Key", key="delete_saved_api_key", on_click=delete_saved_api_key,
                args=(provider,), disabled=running or not saved_api_key, use_container_width=True,
            )
        notice = st.session_state.pop(f"api_key_notice_{provider}", None)
        if notice:
            kind, message = notice
            if kind == "success":
                st.success(message)
            elif kind == "warning":
                st.warning(message)
            else:
                st.error(message)

        model_key = f"model_{provider}"
        if provider == "DeepSeek" and "deepseek_display_migrated" not in st.session_state:
            if st.session_state.get(model_key) == "deepseek-flash":
                st.session_state[model_key] = "DeepSeek-V4.1-Flash"
            st.session_state["deepseek_display_migrated"] = True
        if model_key not in st.session_state:
            st.session_state[model_key] = default_model
        model = st.text_input("模型名称", key=model_key, disabled=running)
        if provider == "DeepSeek":
            st.caption("DeepSeek-V4.1-Flash 的 API 模型 ID 为 deepseek-flash；调用时自动转换。")
        base_url = None
        if provider == "MiMo":
            base_url = st.text_input(
                "接口地址",
                value=os.getenv("MIMO_BASE_URL") or "https://api.xiaomimimo.com/v1",
                help="使用 Token Plan 时，填写控制台提供的专属 OpenAI 兼容地址。",
                key="mimo_base_url",
                disabled=running,
            )
        elif provider == "DeepSeek":
            base_url = st.text_input(
                "接口地址",
                value=os.getenv("DEEPSEEK_BASE_URL") or "https://api.deepseek.com",
                help="DeepSeek 官方 OpenAI 兼容接口地址；可按需填写自定义地址。",
                key="deepseek_base_url",
                disabled=running,
            )

        api_key = api_key_input.strip() or saved_api_key or os.getenv(env_name, "").strip()
        endpoint = base_url.strip() if base_url is not None else None

        st.subheader("决策接口")
        decision_mode = st.selectbox(
            "决策模式", ["跟随主模型", JEV_DECISION_MODE], key="decision_mode",
            disabled=running,
            help="跟随主模型：主模型以 JSON 模式评估；Jev：用 TypeSafe Jev 评估主题。编码、主题和文字反馈仍由主模型生成。",
        )
        jev_api_key = ""
        if decision_mode == JEV_DECISION_MODE:
            saved_jev_api_key = load_saved_api_key("Jev")
            jev_api_key_input = st.text_input(
                "Jev API Key", type="password", key="api_key_Jev", disabled=running,
            )
            if saved_jev_api_key:
                st.caption("已保存 Jev API Key；输入框留空时自动使用。")
            else:
                st.caption("留空时读取本机环境变量 JEV_API_KEY。")
            jev_save_col, jev_delete_col = st.columns(2)
            with jev_save_col:
                st.button(
                    "保存 Jev Key", key="save_jev_api_key", on_click=save_api_key,
                    args=("Jev",), disabled=running, use_container_width=True,
                )
            with jev_delete_col:
                st.button(
                    "删除已存 Key", key="delete_saved_jev_api_key",
                    on_click=delete_saved_api_key, args=("Jev",),
                    disabled=running or not saved_jev_api_key, use_container_width=True,
                )
            jev_notice = st.session_state.pop("api_key_notice_Jev", None)
            if jev_notice:
                kind, message = jev_notice
                if kind == "success":
                    st.success(message)
                elif kind == "warning":
                    st.warning(message)
                else:
                    st.error(message)
            jev_api_key = (
                jev_api_key_input.strip()
                or saved_jev_api_key
                or os.getenv("JEV_API_KEY", "").strip()
            )
            st.caption("可选环境变量：JEV_BASE_URL、JEV_MODEL。")

        st.caption("测试会发送一次简短模型请求，可能产生少量费用。")
        if st.button(
            "测试 API 连接", key="test_api_connection", disabled=running,
            use_container_width=True,
        ):
            error = connection_input_error(
                api_key, model, base_url, endpoint, decision_mode, jev_api_key, env_name,
            )
            if error:
                st.error(error)
            else:
                try:
                    with st.spinner("正在测试模型接口……"):
                        check_api_connection(api_key, model.strip(), endpoint, provider)
                    if decision_mode == JEV_DECISION_MODE:
                        with st.spinner("正在测试决策接口……"):
                            check_jev_connection(
                                jev_api_key,
                                os.getenv("JEV_BASE_URL"),
                            )
                except APIStatusError as exc:
                    st.error(f"连接失败 · HTTP {exc.status_code}，请检查密钥、地址和模型名称。")
                except Exception as exc:
                    st.error("连接失败，请检查网络、接口地址和模型名称。")
                    st.caption(f"错误类型：{type(exc).__name__}")
                else:
                    st.success("连接成功，密钥、地址和模型可用。")

        st.divider()
        st.subheader("切块与评估")
        profile_label = st.selectbox(
            "研究配置", ["通用主题分析", "企业能力与证据链"],
            key="research_profile", disabled=running,
        )
        profile = "enterprise_evidence" if profile_label == "企业能力与证据链" else "generic"
        research_question = st.text_input(
            "研究问题（选填）", key="research_question", disabled=running,
            help="通用配置留空即可；企业配置留空时使用内置研究问题。",
        )
        focus_text = st.text_input(
            "关注点（选填，逗号分隔）", key="focus_areas", disabled=running,
            help="通用配置留空即可；企业配置留空时使用内置关注点。",
        )
        chunk_strategy_label = st.selectbox(
            "切块策略",
            list(CHUNK_STRATEGIES),
            index=1,
            key="chunk_strategy",
            disabled=running,
            help="自动模式会按全文长度计算片段数量，并优先在完整问答、段落、说话人轮次或句末切分。",
        )
        chunk_strategy = CHUNK_STRATEGIES[chunk_strategy_label]
        chunk_size = None
        if chunk_strategy == "manual":
            chunk_size = st.number_input(
                "初始切块上限（字词/块）", min_value=100, max_value=10000,
                value=4000, step=100, key="chunk_size", disabled=running,
                help="中文按字，英文等文本按空格分词；系统仍会优先在自然边界处切分。",
            )
        else:
            strategy_help = {
                "fine": "片段更小，局部编码更细，请求通常更多。",
                "balanced": "在编码细度和 API 请求数量之间保持平衡。",
                "economy": "片段更大，减少 API 请求，适合较长材料。",
            }
            st.caption(strategy_help[chunk_strategy])
        max_workers = st.number_input(
            "并发请求数", min_value=1, max_value=8, value=4, step=1,
            key="max_workers", disabled=running,
            help="仅用于相互独立的片段编码和主题评估请求；遇到接口限流时可调低。",
        )
        max_iterations = st.number_input(
            "最多评估轮次", min_value=1, max_value=10, value=5,
            disabled=running,
        )
        save_intermediate = st.checkbox(
            "保存阶段文件",
            value=False,
            disabled=running,
            help="包含文本片段、编码及逐轮评估，保存在本机 outputs 目录。",
        )
        save_final = st.checkbox(
            "保存本地结果", value=False, key="save_final", disabled=running,
            help="关闭时只在当前服务进程中保留结果；可手动下载 Word 或 JSON。",
        )
        with st.expander("高级设置"):
            confidence_threshold = st.number_input(
                "置信度阈值", min_value=0.05, max_value=1.0, value=0.7, step=0.05,
                key="confidence_threshold", disabled=running,
                help="评估判断的置信度低于该值时，会请求详细反馈并在结果中标记为建议人工复核。",
            )

    st.caption("TAMA · 质性研究工作台")
    st.title("访谈质性分析")
    st.write("提取有原文依据的编码，识别共享意义模式、反例和待核查问题；可按研究问题聚焦。")

    st.divider()
    st.header("1. 输入访谈材料")
    case_id = st.text_input(
        "匿名案例编号（选填）", key="case_id", max_chars=80, disabled=running,
        help="例如 Case-01。编号会进入报告，不会作为文件夹名称。",
    )
    source = st.radio(
        "输入方式", ["粘贴文本", "上传文件"], horizontal=True,
        disabled=running,
    )
    pasted_text = ""
    uploaded_file = None
    transcript = ""
    input_error = None
    if source == "粘贴文本":
        pasted_text = st.text_area(
            "逐字稿内容",
            height=320,
            placeholder="访谈者：请描述一次相关经历。\n受访者：当时我……",
            key="transcript_text",
            disabled=running,
        )
    else:
        uploaded_file = st.file_uploader(
            "选择 UTF-8 TXT 或 Word DOCX 文件", type=["txt", "docx"],
            disabled=running,
        )

    try:
        transcript = read_transcript(source, pasted_text, uploaded_file)
    except UnicodeDecodeError:
        input_error = "文件无法按 UTF-8 读取，请将逐字稿另存为 UTF-8 文本后重试。"
    except ValueError as exc:
        input_error = str(exc)

    if input_error:
        st.error(input_error)
    elif transcript:
        preview, _ = plan_chunks(
            transcript,
            strategy=chunk_strategy,
            manual_chunk_size=int(chunk_size) if chunk_size is not None else None,
        )
        size_unit = "字词" if preview.measurement_unit == "text_units" else "估算模型文本单位"
        with st.container(border=True):
            st.subheader("切块预估")
            st.write(f"已读取 **{preview.total_characters:,} 个字符** · 预计 **{preview.num_chunks} 个片段**")
            st.caption(
                f"每片目标约 {preview.target_size} {size_unit}，上限 {preview.hard_limit}；"
                "优先在问答、说话人轮次、段落或句末切分。"
            )

    st.header("2. 运行分析")
    st.caption("请先对访谈材料去标识化。开始后，文本会发送到所选模型服务；模型生成的判断需要研究者核对。")
    if st.button("开始分析", type="primary", use_container_width=True, key="run_analysis", disabled=running):
        connection_error = connection_input_error(
            api_key, model, base_url, endpoint, decision_mode, jev_api_key, env_name,
        )
        if input_error:
            st.error(input_error)
        elif not transcript:
            st.error("请先粘贴访谈文本或上传 TXT/DOCX 文件。")
        elif connection_error:
            st.error(connection_error)
        else:
            decision_provider = None
            if decision_mode == JEV_DECISION_MODE:
                decision_provider = JevDecisionClient(
                    api_key=jev_api_key,
                    base_url=os.getenv("JEV_BASE_URL") or None,
                    model=os.getenv("JEV_MODEL") or None,
                )

            def run(checkpoint):
                framework = TAMAFramework(
                    api_key=api_key,
                    model=api_model_name(provider, model),
                    base_url=endpoint,
                    chunk_size=int(chunk_size) if chunk_size is not None else None,
                    chunk_strategy=chunk_strategy,
                    max_workers=int(max_workers),
                    max_iterations=int(max_iterations),
                    decision_provider=decision_provider,
                    confidence_threshold=float(confidence_threshold),
                    output_dir=str(ROOT / "outputs"),
                    profile=profile,
                    research_question=research_question,
                    focus_areas=[item.strip() for item in focus_text.replace("，", ",").split(",") if item.strip()] or None,
                )
                return framework.run_analysis(
                    transcript=transcript,
                    save_intermediate=save_intermediate,
                    save_final=save_final,
                    before_model_call=checkpoint,
                    **({"case_id": case_id} if case_id.strip() else {}),
                )

            st.session_state.pop("analysis_result", None)
            try:
                job = JOB_REGISTRY.start(run)
            except RuntimeError:
                st.error("已有分析正在运行，请接回或取消该任务。")
            else:
                st.session_state["analysis_job"] = job
                st.rerun()

    if job is not None and job.snapshot().done:
        snapshot = job.snapshot()
        if snapshot.result is not None:
            st.session_state["analysis_result"] = snapshot.result
        elif snapshot.cancelled:
            st.info("分析已取消；已发出的模型请求可能已计费。")
        else:
            st.error("分析未完成。请检查密钥、接口地址、网络和模型返回内容。")
            st.caption(f"错误类型：{snapshot.error_type}")
            st.caption(f"失败阶段：{snapshot.stage}")
            if snapshot.error_location:
                st.caption(f"错误位置：{snapshot.error_location}")
            if snapshot.http_status:
                st.caption(f"HTTP 状态码：{snapshot.http_status}")
            if snapshot.error_type == "TypeError":
                st.info("程序调用发生类型错误。若刚更新过代码，请重启本地服务后重试。")
    elif running:
        render_active_job()

    render_reflexivity()
    st.divider()
    if "analysis_result" in st.session_state:
        render_result(st.session_state["analysis_result"])
    else:
        st.header("3. 分析结果")
        with st.container(border=True):
            st.subheader("尚无分析结果")
            st.write("完成分析后，这里会显示主题、关联编码和评估结果，并提供 Word 报告与 JSON 数据下载。")

    render_methodology_references()


if __name__ == "__main__":
    main()
